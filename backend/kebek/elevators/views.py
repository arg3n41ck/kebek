import django.utils.timezone as tz
import pandas as pd
import pytz
import time

from io import BytesIO
from datetime import datetime, date
from docx import Document as docx_doc
from weasyprint import HTML

from django.db.models import Sum
from django.http import StreamingHttpResponse
from django.views import generic

from rest_framework import filters, mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound, PermissionDenied, ValidationError
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response

from django_filters.rest_framework import DjangoFilterBackend
from drf_spectacular.utils import extend_schema

from ..notify.utils import send_sms, send_email, create_notification
from ..users.models import CLIENT, ACCOUNTANT, ADMINISTRATOR
from ..users.permissions import IsClient, IsOwner, IsAdministration, IsStaff, IsAdministrationOrClient
from .models import (
    Elevator,
    Accountant,
    Administrator,
    Product,
    Delivery,
    Payment,
    Order,
    OrderItem,
    Document,
    Profit,
    ACTIVE,
    ARCHIVED,
    ACCEPTED,
    WAITING_FOR_PREPAYMENT,
    BILLED,
    PAID,
    PROXY_ADDED,
    FINISHED,
    CANCELLED,
)
from .pagination import ProductsSetPagination
from .serializers import (
    ElevatorFullSerializer,
    ElevatorOrderSerializer,
    ElevatorProfitSerializer,
    ProductCreateSerializer,
    ProductSerializer,
    ProductStatusSerializer,
    DeliveryCreateSerializer,
    DeliverySerializer,
    DeliveryStatusSerializer,
    PaymentCreateSerializer,
    PaymentSerializer,
    PaymentStatusSerializer,
    OrderCreateSerializer,
    OrderPatchSerializer,
    OrderSingleSerializer,
    OrderSerializer,
    OrderProxySerializer,
    OrderStatusSerializer,
    OrderIdsSerializer,
    DocumentCreateSerializer,
    DocumentSerializer,
    StatusIdsSerializer,
    ElevatorGeneralSerializer,
)
from .tasks import revoke_task
from .utils import is_allowed, create_qr


class ElevatorReadOnlyViewSet(viewsets.ReadOnlyModelViewSet):
    """
    Lists and retrieves elevators for all type of users
    """

    permission_classes = [AllowAny]
    serializer_class = ElevatorFullSerializer
    queryset = Elevator.objects.all()

    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['cities']

    def get_serializer_context(self):
        return {'request': self.request}


class ElevatorViewSet(mixins.ListModelMixin, mixins.RetrieveModelMixin, mixins.UpdateModelMixin,
                      viewsets.GenericViewSet):
    """
    Lists, retrieves and edits elevators
    """

    serializer_class = ElevatorFullSerializer
    permission_classes_by_action = {
        'default': [IsAdministration],
        'retrieve': [IsStaff],
        'list': [IsOwner]
    }
    filter_backends = [filters.SearchFilter]
    search_fields = ['title_ru', 'title_kk']

    def get_permissions(self):
        try:
            return [permission() for permission in self.permission_classes_by_action[self.action]]
        except KeyError:
            return [permission() for permission in self.permission_classes_by_action['default']]

    def get_serializer_context(self):
        return {'request': self.request}

    def get_object(self):
        try:
            elevator = Elevator.objects.get(pk=self.kwargs['pk'])
        except Elevator.DoesNotExist:
            raise NotFound()

        is_allowed(self.request.user, elevator)

        return elevator

    def get_queryset(self):
        return Elevator.objects.filter(owner=self.request.user)

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object())

        instance = serializer.save()

        if 'logo' in self.request.FILES:
            instance.logo = self.request.FILES['logo']
            instance.save()

        return instance


class ProductReadOnlyViewSet(viewsets.ReadOnlyModelViewSet):
    """
    Lists and retrieves products for all type of users
    """

    permission_classes = [AllowAny]
    serializer_class = ProductSerializer
    queryset = Product.objects.filter(status=ACTIVE)
    pagination_class = ProductsSetPagination

    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['elevator', 'type', 'elevator__cities']

    @extend_schema(
        responses={200: ProductSerializer},
    )
    @action(
        detail=False,
        methods=['get'],
        url_path='landing',
        url_name='landing',
    )
    def landing(self, request):
        products = Product.objects.filter(status=ACTIVE, display=True)

        page = self.paginate_queryset(products)
        if page is not None:
            serializer = self.serializer_class(page, many=True, context={'request': request})
            return self.get_paginated_response(serializer.data)

        serializer = self.serializer_class(products, many=True, context={'request': request})

        return Response(serializer.data, status=status.HTTP_200_OK)


class ProductViewSet(viewsets.ModelViewSet):
    """
    Lists, filters, retrieves, creates, edits and deletes products
    """

    serializer_class = ProductSerializer
    queryset = Product.objects.all()

    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['elevator', 'type', 'status']

    permission_classes_by_action = {
        'default': [IsAdministration],
        'retrieve': [IsStaff],
        'list': [IsStaff],
        'change_status_bulk': [IsAdministration],
        'change_status': [IsAdministration]
    }

    def get_permissions(self):
        try:
            return [permission() for permission in self.permission_classes_by_action[self.action]]
        except KeyError:
            return [permission() for permission in self.permission_classes_by_action['default']]

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return self.serializer_class
        return ProductCreateSerializer

    def get_object(self):
        try:
            product = Product.objects.get(pk=self.kwargs['pk'])
        except Product.DoesNotExist:
            raise NotFound()

        is_allowed(self.request.user, product.elevator)

        return product

    def get_queryset(self):
        user = self.request.user

        if user.user_role == ACCOUNTANT:
            return Product.objects.filter(elevator__elevator_accountant__accountant=user)
        elif user.user_role == ADMINISTRATOR:
            return Product.objects.filter(elevator__elevator_administrator__administrator=user)
        else:
            return Product.objects.filter(elevator__owner=user)

    def create(self, request, *args, **kwargs):
        serializer = ProductCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_create(serializer)
        instance_serializer = ProductSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = ProductCreateSerializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_update(serializer)
        instance_serializer = ProductSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        user = self.request.user

        try:
            elevator = Elevator.objects.get(pk=self.request.data['elevator'])
        except Elevator.DoesNotExist:
            raise NotFound('Elevator not found')

        is_allowed(user, elevator)

        if user.user_role == ADMINISTRATOR:
            instance = serializer.save(elevator=elevator)
        else:
            instance = serializer.save()

        if 'image' in self.request.FILES:
            instance.image = self.request.FILES['image']
            instance.save()

        return instance

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object().elevator)

        instance = serializer.save()

        if 'image' in self.request.FILES:
            instance.image = self.request.FILES['image']
            instance.save()

        return instance

    def perform_destroy(self, instance):
        is_allowed(self.request.user, instance.elevator)

        instance.status = ARCHIVED
        instance.save()

    @extend_schema(
        request=StatusIdsSerializer,
        responses={200: ProductSerializer},
    )
    @action(
        detail=False,
        methods=['post'],
        url_path='status/bulk',
        url_name='status/bulk'
    )
    def change_status_bulk(self, request):
        for product_id in request.data['ids']:
            try:
                product = Product.objects.get(id=product_id)
                product.status = request.data['status']
                product.save()
            except Product.DoesNotExist:
                pass

        page = self.paginate_queryset(self.get_queryset())
        if page is not None:
            serializer = ProductSerializer(page, context={'request': request}, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = ProductSerializer(self.get_queryset(), context={'request': request}, many=True)

        return Response(serializer.data, status=status.HTTP_200_OK)

    @extend_schema(
        request=ProductStatusSerializer,
        responses={200: ProductSerializer},
    )
    @action(
        detail=True,
        methods=['patch'],
        url_path='status',
        url_name='status',
    )
    def change_status(self, request, pk):
        try:
            instance = Product.objects.get(id=pk)
        except Product.DoesNotExist:
            raise NotFound()

        instance.status = request.data['status']
        instance.save()

        return Response(ProductSerializer(instance, context={'request': request}).data, status=status.HTTP_200_OK)


class DeliveryViewSet(viewsets.ModelViewSet):
    """
    Lists, filters, retrieves, creates, edits and deletes deliveries
    """

    serializer_class = DeliverySerializer
    queryset = Delivery.objects.all()

    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['elevator', 'type', 'status']

    permission_classes_by_action = {
        'default': [IsAdministration],
        'retrieve': [IsStaff],
        'list': [IsStaff],
        'change_status_bulk': [IsAdministration],
        'change_status': [IsAdministration]
    }

    def get_permissions(self):
        try:
            return [permission() for permission in self.permission_classes_by_action[self.action]]
        except KeyError:
            return [permission() for permission in self.permission_classes_by_action['default']]

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return self.serializer_class
        return DeliveryCreateSerializer

    def get_object(self):
        try:
            delivery = Delivery.objects.get(pk=self.kwargs['pk'])
        except Delivery.DoesNotExist:
            raise NotFound()

        is_allowed(self.request.user, delivery.elevator)

        return delivery

    def get_queryset(self):
        user = self.request.user

        if user.user_role == ACCOUNTANT:
            return Delivery.objects.filter(elevator__elevator_accountant__accountant=user)
        elif user.user_role == ADMINISTRATOR:
            return Delivery.objects.filter(elevator__elevator_administrator__administrator=user)
        else:
            return Delivery.objects.filter(elevator__owner=user)

    def create(self, request, *args, **kwargs):
        serializer = DeliveryCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_create(serializer)
        instance_serializer = DeliverySerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = DeliveryCreateSerializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_update(serializer)
        instance_serializer = DeliverySerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        user = self.request.user

        try:
            elevator = Elevator.objects.get(pk=self.request.data['elevator'])
        except Elevator.DoesNotExist:
            raise NotFound('Elevator not found')

        is_allowed(user, elevator)

        if user.user_role == ADMINISTRATOR:
            instance = serializer.save(elevator=elevator)
        else:
            instance = serializer.save()

        return instance

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object().elevator)

        instance = serializer.save()

        return instance

    def perform_destroy(self, instance):
        is_allowed(self.request.user, instance.elevator)

        instance.status = ARCHIVED
        instance.save()

    @extend_schema(
        request=StatusIdsSerializer,
        responses={200: DeliverySerializer},
    )
    @action(
        detail=False,
        methods=['post'],
        url_path='status/bulk',
        url_name='status/bulk'
    )
    def change_status_bulk(self, request):
        for delivery_id in request.data['ids']:
            try:
                delivery = Delivery.objects.get(id=delivery_id)
                delivery.status = request.data['status']
                delivery.save()
            except Delivery.DoesNotExist:
                pass

        page = self.paginate_queryset(self.get_queryset())
        if page is not None:
            serializer = DeliverySerializer(page, context={'request': request}, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = DeliverySerializer(self.get_queryset(), context={'request': request}, many=True)

        return Response(serializer.data, status=status.HTTP_200_OK)

    @extend_schema(
        request=DeliveryStatusSerializer,
        responses={200: DeliverySerializer},
    )
    @action(
        detail=True,
        methods=['patch'],
        url_path='status',
        url_name='status',
    )
    def change_status(self, request, pk):
        try:
            instance = Delivery.objects.get(id=pk)
        except Delivery.DoesNotExist:
            raise NotFound()

        instance.status = request.data['status']
        instance.save()

        return Response(DeliverySerializer(instance, context={'request': request}).data, status=status.HTTP_200_OK)


class PaymentViewSet(viewsets.ModelViewSet):
    """
    Lists, filters, retrieves, creates, edits and deletes payments
    """

    serializer_class = PaymentSerializer
    queryset = Payment.objects.all()

    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['elevator', 'type', 'status']

    permission_classes_by_action = {
        'default': [IsAdministration],
        'retrieve': [IsStaff],
        'list': [IsStaff],
        'change_status_bulk': [IsAdministration],
        'change_status': [IsAdministration]
    }

    def get_permissions(self):
        try:
            return [permission() for permission in self.permission_classes_by_action[self.action]]
        except KeyError:
            return [permission() for permission in self.permission_classes_by_action['default']]

    def get_serializer_class(self):
        if self.action in ['list', 'retrieve']:
            return self.serializer_class
        return PaymentCreateSerializer

    def get_object(self):
        try:
            payment = Payment.objects.get(pk=self.kwargs['pk'])
        except Payment.DoesNotExist:
            raise NotFound()

        is_allowed(self.request.user, payment.elevator)

        return payment

    def get_queryset(self):
        user = self.request.user

        if user.user_role == ACCOUNTANT:
            return Payment.objects.filter(elevator__elevator_accountant__accountant=user)
        elif user.user_role == ADMINISTRATOR:
            return Payment.objects.filter(elevator__elevator_administrator__administrator=user)
        else:
            return Payment.objects.filter(elevator__owner=user)

    def create(self, request, *args, **kwargs):
        serializer = PaymentCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_create(serializer)
        instance_serializer = PaymentSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = PaymentCreateSerializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_update(serializer)
        instance_serializer = PaymentSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        user = self.request.user

        try:
            elevator = Elevator.objects.get(pk=self.request.data['elevator'])
        except Elevator.DoesNotExist:
            raise NotFound('Elevator not found')

        is_allowed(user, elevator)

        if user.user_role == ADMINISTRATOR:
            instance = serializer.save(elevator=elevator)
        else:
            instance = serializer.save()

        return instance

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object().elevator)

        instance = serializer.save()

        return instance

    def perform_destroy(self, instance):
        is_allowed(self.request.user, instance.elevator)

        instance.status = ARCHIVED
        instance.save()

    @extend_schema(
        request=StatusIdsSerializer,
        responses={200: PaymentSerializer},
    )
    @action(
        detail=False,
        methods=['post'],
        url_path='status/bulk',
        url_name='status/bulk'
    )
    def change_status_bulk(self, request):
        for payment_id in request.data['ids']:
            try:
                payment = Payment.objects.get(id=payment_id)
                payment.status = request.data['status']
                payment.save()
            except Payment.DoesNotExist:
                pass

        page = self.paginate_queryset(self.get_queryset())
        if page is not None:
            serializer = PaymentSerializer(page, context={'request': request}, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = PaymentSerializer(self.get_queryset(), context={'request': request}, many=True)

        return Response(serializer.data, status=status.HTTP_200_OK)

    @extend_schema(
        request=PaymentStatusSerializer,
        responses={200: PaymentSerializer},
    )
    @action(
        detail=True,
        methods=['patch'],
        url_path='status',
        url_name='status',
    )
    def change_status(self, request, pk):
        try:
            instance = Payment.objects.get(id=pk)
        except Payment.DoesNotExist:
            raise NotFound()

        instance.status = request.data['status']
        instance.save()

        return Response(PaymentSerializer(instance, context={'request': request}).data, status=status.HTTP_200_OK)


class OrderViewSet(viewsets.ModelViewSet):
    """
    Lists, filters, retrieves, creates, edits and deletes orders
    """

    serializer_class = OrderSerializer
    queryset = Order.objects.all()

    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    filterset_fields = ['elevator', 'client', 'delivery', 'payment', 'status']
    search_fields = ['number']

    permission_classes_by_action = {
        'default': [IsAdministration],
        'retrieve': [IsAuthenticated],
        'list': [IsAuthenticated],
        'create': [IsAdministrationOrClient],
        'latest': [IsClient],
        'change_status': [IsStaff],
        'proxy': [IsClient],
        'report_xlsx': [IsAdministration],
        'report_docx': [IsAdministration],
        'report_pdf': [IsAdministration],
        'delete_orders_bulk': [IsAdministration],
    }

    def get_permissions(self):
        try:
            return [permission() for permission in self.permission_classes_by_action[self.action]]
        except KeyError:
            return [permission() for permission in self.permission_classes_by_action['default']]

    def get_serializer_class(self):
        if self.action == 'retrieve':
            return OrderSingleSerializer
        elif self.action == 'update':
            return OrderPatchSerializer
        elif self.action == 'create':
            return OrderCreateSerializer
        else:
            return self.serializer_class

    def get_serializer_context(self):
        return {'request': self.request}

    def get_object(self):
        user = self.request.user

        try:
            order = Order.objects.get(pk=self.kwargs['pk'])
        except Order.DoesNotExist:
            raise NotFound()

        if user.user_role == CLIENT:
            if order.client != user:
                raise PermissionDenied()
        else:
            is_allowed(user, order.elevator)

        return order

    def get_queryset(self):
        user = self.request.user
        product = self.request.query_params.get('product')

        if user.user_role == CLIENT:
            queryset = Order.objects.filter(client=user)
        elif user.user_role == ACCOUNTANT:
            queryset = Order.objects.filter(elevator__elevator_accountant__accountant=user)
        elif user.user_role == ADMINISTRATOR:
            queryset = Order.objects.filter(elevator__elevator_administrator__administrator=user)
        else:
            queryset = Order.objects.filter(elevator__owner=user)

        if product:
            queryset = queryset.filter(items__product__type=product)

        return queryset

    def create(self, request, *args, **kwargs):
        serializer = OrderCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_create(serializer)
        instance_serializer = OrderSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = OrderPatchSerializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_update(serializer)
        instance_serializer = OrderSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_200_OK)

    def perform_create(self, serializer):
        user = self.request.user
        items = self.request.data['products']
        payment = self.request.data['payment']
        delivery = self.request.data['delivery']

        try:
            elevator = Elevator.objects.get(pk=self.request.data['elevator'])
        except Elevator.DoesNotExist:
            raise NotFound('Elevator not found')

        if user.user_role != CLIENT:
            is_allowed(user, elevator)

        try:
            Payment.objects.get(elevator=elevator, pk=payment, status=ACTIVE)
        except Payment.DoesNotExist:
            raise NotFound('Payment not found')

        try:
            Delivery.objects.get(elevator=elevator, pk=delivery, status=ACTIVE)
        except Delivery.DoesNotExist:
            raise NotFound('Delivery not found')

        products = Product.objects.filter(id__in=[item['product'] for item in items], elevator=elevator, status=ACTIVE)

        if products.count() != len(items):
            raise NotFound('At least one product is not found in elevator')

        if user.user_role == CLIENT:
            instance = serializer.save(client=user)
        elif user.user_role == ADMINISTRATOR:
            instance = serializer.save(elevator=elevator)
        else:
            instance = serializer.save()

        for item in items:
            product = Product.objects.get(pk=item['product'])
            OrderItem.objects.create(
                order=instance,
                product=product,
                amount=item['amount'],
                product_payment=item['product_payment']
            )

        instance.schedule_payment_expiration()

        content = f'Заказ №{instance.number} создан и принят в обработку'
        title = 'Создан'

        create_notification(user, instance, ACCEPTED, title, content)

        if instance.client.notifications_sms:
            send_sms(instance.client.username, content)
        if instance.client.notifications_email and instance.client.email:
            send_email(instance.client.email, title, content)

        return instance

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object().elevator)

        items = self.request.data['products']

        try:
            elevator = Elevator.objects.get(pk=self.request.data['elevator'])
        except Elevator.DoesNotExist:
            raise NotFound('Elevator not found')

        if 'payment' in self.request.data:
            try:
                Payment.objects.get(elevator=elevator, pk=self.request.data['payment'], status=ACTIVE)
            except Payment.DoesNotExist:
                raise NotFound('Payment not found')

        if 'delivery' in self.request.data:
            try:
                Delivery.objects.get(elevator=elevator, pk=self.request.data['delivery'], status=ACTIVE)
            except Delivery.DoesNotExist:
                raise NotFound('Delivery not found')

        products = Product.objects.filter(id__in=[item['product'] for item in items], elevator=elevator, status=ACTIVE)

        if products.count() != len(items):
            raise NotFound('At least one product is not found in elevator')

        OrderItem.objects.filter(order=self.get_object()).delete()

        for item in items:
            product = Product.objects.get(pk=item['product'])
            OrderItem.objects.create(
                order=self.get_object(),
                product=product,
                amount=item['amount'],
                product_payment=item['product_payment']
            )

        instance = serializer.save()

        return instance

    def perform_destroy(self, instance):
        is_allowed(self.request.user, instance.elevator)

        revoke_task(instance, 'elevators.payment_expired')

        instance.delete()

    @extend_schema(
        responses={200: OrderSerializer},
    )
    @action(
        detail=False,
        methods=['get'],
        url_path='latest',
        url_name='latest',
    )
    def latest(self, request):
        orders = Order.objects.filter(client=request.user)[:3]

        page = self.paginate_queryset(orders)
        if page is not None:
            serializer = OrderSerializer(page, many=True, context={'request': request})
            return self.get_paginated_response(serializer.data)

        serializer = OrderSerializer(orders, many=True, context={'request': request})

        return Response(serializer.data, status=status.HTTP_200_OK)

    @extend_schema(
        request=OrderStatusSerializer,
        responses={200: OrderSerializer},
    )
    @action(
        detail=True,
        methods=['patch'],
        url_path='status',
        url_name='status',
    )
    def change_status(self, request, pk):
        user = request.user
        order_status = request.data['status']

        if user.user_role == ADMINISTRATOR:
            try:
                instance = Order.objects.get(elevator__elevator_administrator__administrator=user, id=pk)
            except Order.DoesNotExist:
                raise NotFound()
        else:
            try:
                instance = Order.objects.get(elevator__owner=user, id=pk)
            except Order.DoesNotExist:
                raise NotFound()

        if order_status not in [PAID, FINISHED, CANCELLED] and instance.status == PAID:
            try:
                Profit.objects.get(elevator=instance.elevator, order=instance).delete()
            except Profit.DoesNotExist:
                pass
        elif order_status == PAID and instance.status not in [PAID, FINISHED, CANCELLED]:
            payment = OrderItem.objects.filter(order=instance).aggregate(Sum('product_payment'))['product_payment__sum']
            try:
                profit = Profit.objects.get(elevator=instance.elevator, order=instance)
                profit.profit = payment
                profit.save()
            except Profit.DoesNotExist:
                Profit.objects.create(elevator=instance.elevator, order=instance, profit=payment)

        if order_status == BILLED and instance.status != BILLED:
            content = f'Выставлен счет для оплаты заказа №{instance.number}. Скачайте его в личном кабинете kebek.kz'
            title = 'Выставлен счет'

            create_notification(instance.client, instance, BILLED, title, content)

            if instance.client.notifications_sms:
                send_sms(instance.client.username, content)
            if instance.client.notifications_email and instance.client.email:
                send_email(instance.client.email, title, content)

        elif order_status == PAID and instance.status != PAID:
            create_qr(instance)

            content = f'Заказ №{instance.number} оплачен. Спасибо за использование системы Kebek!'
            title = 'Оплачен'

            create_notification(instance.client, instance, PAID, title, content)

            if instance.client.notifications_sms:
                send_sms(instance.client.username, content)
            if instance.client.notifications_email and instance.client.email:
                send_email(instance.client.email, title, content)

        elif order_status == FINISHED and instance.status != FINISHED:
            items = OrderItem.objects.filter(order=instance)
            products = Product.objects.filter(id__in=items)

            for item in items:
                for product in products:
                    if item.product == product:
                        product.residue -= item.amount
                        product.save()

            content = f'Заказ №{instance.number} получен'
            title = 'Завершен'

            create_notification(instance.client, instance, FINISHED, title, content)

            if instance.client.notifications_sms:
                send_sms(instance.client.username, content)
            if instance.client.notifications_email and instance.client.email:
                send_email(instance.client.email, title, content)

        elif order_status == CANCELLED and instance.status != CANCELLED:
            content = f'Заказ №{instance.number} отменен из-за отсутствия оплаты'
            title = 'Отменен'

            create_notification(instance.client, instance, CANCELLED, title, content)

            if instance.client.notifications_sms:
                send_sms(instance.client.username, content)
            if instance.client.notifications_email and instance.client.email:
                send_email(instance.client.email, title, content)

        instance.status = order_status
        instance.save()

        serializer = OrderSerializer(instance, context={'request': request})

        return Response(serializer.data, status=status.HTTP_200_OK)

    @extend_schema(
        request=OrderProxySerializer,
        responses={200: OrderSerializer},
    )
    @action(
        detail=True,
        methods=['patch', 'delete'],
        url_path='proxy',
        url_name='proxy',
    )
    def proxy(self, request, pk):
        try:
            instance = Order.objects.get(client=request.user, id=pk)
        except Order.DoesNotExist:
            raise NotFound()

        if instance.status in [FINISHED, CANCELLED]:
            raise PermissionDenied()

        if request.method == 'PATCH':
            pre_serializer = OrderProxySerializer(instance, data=request.data)
            if pre_serializer.is_valid():
                instance = pre_serializer.save()

                if instance.status != PROXY_ADDED:
                    content = f'К заказу №{instance.number} было добавлено доверенное лицо'
                    title = 'Добавлено доверенное лицо'

                    accountants = Accountant.objects.filter(elevator=instance.elevator)
                    administrators = Administrator.objects.filter(elevator=instance.elevator)

                    create_notification(instance.elevator.owner, instance, PROXY_ADDED, title, content)

                    if instance.elevator.owner.notifications_sms:
                        send_sms(instance.elevator.owner.username, content)
                    if instance.elevator.owner.notifications_email and instance.elevator.owner.email:
                        send_email(instance.elevator.owner.email, title, content)

                    for administrator in administrators:
                        create_notification(administrator.administrator, instance, PROXY_ADDED, title, content)

                        if administrator.administrator.notifications_sms:
                            send_sms(administrator.administrator.username, content)
                        if administrator.administrator.notifications_email and administrator.administrator.email:
                            send_email(administrator.administrator.email, title, content)

                    for accountant in accountants:
                        create_notification(accountant.accountant, instance, PROXY_ADDED, title, content)

                        if accountant.accountant.notifications_sms:
                            send_sms(accountant.accountant.username, content)
                        if accountant.accountant.notifications_email and accountant.accountant.email:
                            send_email(accountant.accountant.email, title, content)

                serializer = OrderSerializer(instance, context={'request': request})
                return Response(serializer.data, status=status.HTTP_200_OK)

            return Response(pre_serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        elif request.method == 'DELETE':
            instance.proxy_fullname = None
            instance.proxy_number = None
            instance.proxy_start_date = None
            instance.proxy_end_date = None
            instance.save()

            return Response(status=status.HTTP_204_NO_CONTENT)

        else:
            return Response(status=status.HTTP_405_METHOD_NOT_ALLOWED)

    @extend_schema(
        request=OrderIdsSerializer,
        responses={200: OrderSerializer},
    )
    @action(
        detail=False,
        methods=['post'],
        url_path='bulk',
        url_name='bulk'
    )
    def delete_orders_bulk(self, request):
        for order_id in request.data['ids']:
            try:
                order = Order.objects.get(id=order_id)
                order.delete()
            except Product.DoesNotExist:
                pass

        page = self.paginate_queryset(self.get_queryset())
        if page is not None:
            serializer = OrderSerializer(page, context={'request': request}, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = OrderSerializer(self.get_queryset(), context={'request': request}, many=True)

        return Response(serializer.data, status=status.HTTP_200_OK)


class OrderReportViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    """
    Downloads orders reports
    """

    permission_classes = [IsAdministration]
    serializer_class = OrderSerializer
    queryset = Order.objects.all()

    def list(self, request, *args, **kwargs):
        user = request.user
        elevator = request.query_params.get('elevator')
        client = request.query_params.get('client')
        product = request.query_params.get('product')
        delivery = request.query_params.get('delivery')
        payment = request.query_params.get('payment')
        order_status = request.query_params.get('status')
        search = request.query_params.get('search')
        report_type = request.query_params.get('report_type')

        queryset = []

        if user.user_role == ACCOUNTANT:
            qs = Order.objects.filter(elevator__elevator_accountant__accountant=user)
        elif user.user_role == ADMINISTRATOR:
            qs = Order.objects.filter(elevator__elevator_administrator__administrator=user)
        else:
            qs = Order.objects.filter(elevator__owner=user)

        if elevator:
            qs = qs.filter(elevator=elevator)
        if client:
            qs = qs.filter(client=client)
        if product:
            qs = qs.filter(items__product__type=product)
        if delivery:
            qs = qs.filter(delivery=delivery)
        if payment:
            qs = qs.filter(payment=payment)
        if order_status:
            qs = qs.filter(status=order_status)
        if search:
            qs = qs.filter(number__icontains=search)

        for query in qs:
            data = {}
            items = OrderItem.objects.filter(order=query.id)

            if query.status == ACCEPTED:
                order_status = 'Создан'
            elif query.status == WAITING_FOR_PREPAYMENT:
                order_status = 'Ожидает оплаты'
            elif query.status == BILLED:
                order_status = 'Выставлен счет'
            elif query.status == PAID:
                order_status = 'Оплачен'
            elif query.status == FINISHED:
                order_status = 'Завершен'
            elif query.status == CANCELLED:
                order_status = 'Отменен'
            else:
                order_status = ''

            data['Номер'] = query.number
            data['Товары'] = '; '.join([item.product.type.title_ru + ', ' + str(item.amount) for item in items])
            data['Стоимость'] = items.aggregate(Sum('product_payment'))['product_payment__sum']
            data['Тип оплаты'] = query.payment.type.title_ru
            data['Тип доставки'] = query.delivery.type.title_ru
            data['Поставщик'] = query.elevator.title_ru
            data['Клиент'] = query.client.first_name + ' (' + str(query.client.phone_number) + ')'
            data['Дата создания'] = query.created_at.strftime('%Y-%m-%d %H:%M')
            data['Статус'] = order_status

            queryset.append(data)

        columns = [
            'Номер',
            'Товары',
            'Стоимость',
            'Тип оплаты',
            'Тип доставки',
            'Поставщик',
            'Клиент',
            'Дата создания',
            'Статус'
        ]

        df = pd.DataFrame(queryset, columns=columns)
        df.index += 1

        if report_type == 'xlsx':
            output = BytesIO()

            writer = pd.ExcelWriter(output, engine='xlsxwriter')
            df.to_excel(writer, encoding='utf-8')
            writer.save()

            output.seek(0)

            response = StreamingHttpResponse(
                output,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response[
                'Content-Disposition'] = f'attachment; filename=kebek_report_{datetime.timestamp(datetime.now())}.xlsx'
            response['Content-Encoding'] = 'UTF-8'

            return response

        elif report_type == 'docx':
            doc = docx_doc()

            t = doc.add_table(df.shape[0] + 1, df.shape[1])
            t.style = 'TableGrid'

            for j in range(df.shape[-1]):
                t.cell(0, j).text = df.columns[j]

            for i in range(df.shape[0]):
                for j in range(df.shape[-1]):
                    t.cell(i + 1, j).text = str(df.values[i, j])

            output = BytesIO()
            doc.save(output)
            output.seek(0)

            response = StreamingHttpResponse(
                output,
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = f'attachment; filename=kebek_report_{datetime.timestamp(datetime.now())}.docx'
            response['Content-Encoding'] = 'UTF-8'

            return response

        elif report_type == 'pdf':
            html_pre = '''
                    <!DOCTYPE html>
                    <html>
                    <head>
                        <style>
                            @page {
                                size: A4 landscape;
                                margin: 0.5em;
                            }
                            table {
                                font-size: 10pt;
                                border-collapse: collapse;
                                border: 1px solid black;
                                color: black;
                                margin: auto;
                            }
                            tr:nth-child(odd) {
                                background: #eee;
                            }
                            tr:first-child {
                                background: white;
                            }
                            td, th {
                                border: 1px solid black;
                                padding: 0.5em;
                            }
                            th {
                                text-align: center;
                            }
                        </style>
                    </head>
                    '''

            html_post = '</body></html>'

            html = HTML(string=html_pre + str(df.to_html()) + html_post)

            output = BytesIO()
            html.write_pdf(output)
            output.seek(0)

            response = StreamingHttpResponse(output, content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename=kebek_report_{datetime.timestamp(datetime.now())}.pdf'
            response['Content-Encoding'] = 'UTF-8'

            return response

        else:
            raise NotFound()


class DocumentViewSet(mixins.CreateModelMixin, mixins.UpdateModelMixin, mixins.DestroyModelMixin, viewsets.GenericViewSet):
    """
    Creates, edits and deletes order documents
    """

    permission_classes = [IsAdministration]
    serializer_class = DocumentCreateSerializer
    queryset = Document.objects.all()

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_create(serializer)
        instance_serializer = DocumentSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        instance = self.perform_update(serializer)
        instance_serializer = DocumentSerializer(instance, context={'request': request})
        return Response(instance_serializer.data, status=status.HTTP_200_OK)

    def perform_update(self, serializer):
        is_allowed(self.request.user, self.get_object().order.elevator)

        instance = serializer.save()

        return instance

    def perform_destroy(self, instance):
        is_allowed(self.request.user, instance.order.elevator)

        instance.delete()


class DashboardGeneralViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    """
    Lists general info for dashboard
    """

    permission_classes = [IsAdministration]
    serializer_class = ElevatorGeneralSerializer
    queryset = Elevator.objects.all()

    def list(self, request, *args, **kwargs):
        user = request.user

        if user.user_role == ADMINISTRATOR:
            try:
                administrator_elevators = Administrator.objects.filter(administrator=user).values_list('elevator', flat=True)
                elevators = Elevator.objects.filter(id__in=administrator_elevators)
            except Administrator.DoesNotExist:
                raise NotFound('Administrator does not have an assigned elevator')
        else:
            elevators = Elevator.objects.filter(owner=user)

        today = date.today()

        orders = Order.objects.filter(elevator__in=elevators, created_at__contains=today).count()
        orders_FD = Order.objects.filter(elevator__in=elevators, status=FINISHED, updated_at__contains=today).count()
        profit = Profit.objects.filter(elevator__in=elevators, created_at__contains=today).aggregate(Sum('profit'))['profit__sum']

        return Response({
            'orders': orders or 0,
            'orders_FD': orders_FD or 0,
            'profit': profit or 0
        }, status=status.HTTP_200_OK)


class DashboardOrdersViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    """
    Lists orders for dashboard
    """

    permission_classes = [IsAdministration]
    serializer_class = ElevatorOrderSerializer

    def get_queryset(self):
        user = self.request.user

        if user.user_role == ADMINISTRATOR:
            try:
                administrator_elevators = Administrator.objects.filter(administrator=user).values_list('elevator', flat=True)
                return Elevator.objects.filter(id__in=administrator_elevators)
            except Administrator.DoesNotExist:
                raise NotFound('Administrator does not have an assigned elevator')
        else:
            return Elevator.objects.filter(owner=user)


class DashboardProfitViewSet(mixins.ListModelMixin, viewsets.GenericViewSet):
    """
    Lists profits for dashboard
    """

    permission_classes = [IsAdministration]
    serializer_class = ElevatorProfitSerializer
    queryset = Elevator.objects.all()

    def list(self, request, *args, **kwargs):
        user = request.user

        first_day = tz.make_aware(datetime(datetime.now().year, 1, 1), pytz.timezone('Asia/Almaty')).date()
        now = tz.make_aware(datetime.now(), pytz.timezone('Asia/Almaty')).date()
        start_date = self.request.query_params.get('start_date') or str(first_day)
        end_date = self.request.query_params.get('end_date') or str(now)

        try:
            time.strptime(start_date, '%Y-%m-%d')
            time.strptime(end_date, '%Y-%m-%d')
        except ValueError:
            raise ValidationError({'detail': 'start_date & end_date values must have format YYYY-MM-DD'})

        data = {}
        elevators_list = []

        if user.user_role == ADMINISTRATOR:
            try:
                administrator_elevators = Administrator.objects.filter(administrator=user).values_list('elevator', flat=True)
                elevators = Elevator.objects.filter(id__in=administrator_elevators)
            except Administrator.DoesNotExist:
                raise NotFound('Administrator does not have an assigned elevator')
        else:
            elevators = Elevator.objects.filter(owner=user)

        for elevator in elevators:
            profits_data = []
            profits = Profit.objects.filter(elevator=elevator, created_at__gte=start_date, created_at__lte=end_date)
            dates = profits.values_list('created_at', flat=True).distinct()

            for date in dates:
                profit = profits.filter(created_at=date).aggregate(Sum('profit'))['profit__sum']
                profit_data = {
                    'created_at': date,
                    'profit': profit or 0
                }
                profits_data.append(profit_data)

            elevators_list.append(
                {
                    'id': elevator.id,
                    'title_ru': elevator.title_ru,
                    'title_kk': elevator.title_kk,
                    'profits': profits_data
                }
            )

        data['count'] = elevators.count()
        data['results'] = elevators_list

        return Response(data, status=status.HTTP_200_OK)


class PassDetailView(generic.DetailView):
    """
    Retrieves QR pass
    """

    model = Order
    template_name = 'elevators/qr.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['items'] = OrderItem.objects.filter(order=self.get_object())
        return context
